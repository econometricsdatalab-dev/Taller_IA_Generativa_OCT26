# -*- coding: utf-8 -*-
"""
Base editable de informe técnico de análisis datos con identidad visual EDL.
Genera Base_Informe_EDL.docx (A4, márgenes 2,5 cm, Arial 11, estilos reales,
portada con banda negra #090909 + logo, encabezado y paginación solo en cuerpo).
Requiere: pip install python-docx
Uso:    python generar_base_informe_edl.py [ruta/al/logo_EDL_oficial.png]
"""
import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- constantes
TEXTO   = RGBColor(0x22, 0x22, 0x22)   # texto general
ACENTO  = RGBColor(0xF0, 0x4A, 0x1D)   # acento EDL
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)
BANDA   = "090909"                     # banda negra portada
SALIDA  = "Base_Informe_EDL.docx"
ANCHO_UTIL_CM = 16.0                   # 21,0 - 2,5 - 2,5

# ---------------------------------------------------------------- utilidades XML
def _arial(style):
    style.font.name = "Arial"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:cs"), "Arial")
    rfonts.set(qn("w:eastAsia"), "Arial")

def _shd_celda(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

def _margenes_celda(tabla, top=280, bottom=280, left=140, right=140):
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:%s" % tag)
        el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")
        mar.append(el)
    tabla._tbl.tblPr.insert_element_before(mar, "w:tblLook")

def _borde_inferior(p, color="F04A1D", sz=8):
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), str(sz))
    bot.set(qn("w:space"), "4"); bot.set(qn("w:color"), color)
    pbdr.append(bot); ppr.append(pbdr)

def _campo(p, instr):
    """Inserta un campo de Word (PAGE, NUMPAGES...) en el párrafo."""
    r1 = p.add_run()
    fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin")
    it  = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "separate")
    r1._r.append(fc1); r1._r.append(it); r1._r.append(fc2)
    p.add_run("1")  # resultado provisional hasta actualizar campos en Word
    r3 = p.add_run()
    fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end")
    r3._r.append(fc3)

def _anchos(tabla, anchos_cm):
    tabla.autofit = False
    for fila in tabla.rows:
        for i, c in enumerate(fila.cells):
            c.width = Cm(anchos_cm[i])

# ---------------------------------------------------------------- estilos
def configurar_estilos(doc):
    st = doc.styles

    n = st["Normal"]
    _arial(n); n.font.size = Pt(11); n.font.color.rgb = TEXTO
    n.paragraph_format.line_spacing = 1.15
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.space_before = Pt(0)

    # Jerarquía: reutiliza los estilos de título nativos (conservan nivel de
    # esquema para el panel de navegación) y los renombra a Título 1 / Título 2.
    t1 = st["Heading 1"]; t1.name = "Título 1"
    _arial(t1); t1.font.size = Pt(16); t1.font.bold = True; t1.font.color.rgb = ACENTO
    t1.paragraph_format.space_before = Pt(18); t1.paragraph_format.space_after = Pt(6)
    t1.paragraph_format.keep_with_next = True; t1.paragraph_format.line_spacing = 1.15

    t2 = st["Heading 2"]; t2.name = "Título 2"
    _arial(t2); t2.font.size = Pt(13); t2.font.bold = True; t2.font.color.rgb = TEXTO
    t2.paragraph_format.space_before = Pt(12); t2.paragraph_format.space_after = Pt(6)
    t2.paragraph_format.keep_with_next = True; t2.paragraph_format.line_spacing = 1.15

    ley = st.add_style("Leyenda", WD_STYLE_TYPE.PARAGRAPH); ley.base_style = n
    ley.font.size = Pt(9); ley.font.italic = True; ley.font.color.rgb = TEXTO
    ley.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ley.paragraph_format.space_before = Pt(3); ley.paragraph_format.space_after = Pt(12)
    ley.paragraph_format.line_spacing = 1.15

    tab = st.add_style("Tabla", WD_STYLE_TYPE.PARAGRAPH); tab.base_style = n
    tab.font.size = Pt(10); tab.font.color.rgb = TEXTO
    tab.paragraph_format.line_spacing = 1.0
    tab.paragraph_format.space_after = Pt(0); tab.paragraph_format.space_before = Pt(0)

    pt = st.add_style("Portada Título", WD_STYLE_TYPE.PARAGRAPH); pt.base_style = n
    pt.font.size = Pt(26); pt.font.bold = True; pt.font.color.rgb = TEXTO
    pt.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pt.paragraph_format.space_after = Pt(6)

    ps = st.add_style("Portada Subtítulo", WD_STYLE_TYPE.PARAGRAPH); ps.base_style = n
    ps.font.size = Pt(14); ps.font.color.rgb = ACENTO
    ps.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ps.paragraph_format.space_after = Pt(30)

    pm = st.add_style("Portada Metadatos", WD_STYLE_TYPE.PARAGRAPH); pm.base_style = n
    pm.font.size = Pt(11); pm.font.color.rgb = TEXTO
    pm.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pm.paragraph_format.space_after = Pt(4)

# ---------------------------------------------------------------- página
def configurar_pagina(sec):
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.25); sec.footer_distance = Cm(1.25)

# ---------------------------------------------------------------- portada
def portada(doc, logo):
    banda = doc.add_table(rows=1, cols=1)
    celda = banda.cell(0, 0)
    _shd_celda(celda, BANDA)          # sombreado antes que vAlign/width (orden OOXML)
    _margenes_celda(banda)
    celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _anchos(banda, [ANCHO_UTIL_CM])
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    if logo and os.path.exists(logo):
        p.add_run().add_picture(logo, width=Cm(12))   # solo ancho => proporción intacta
    else:
        r = p.add_run("[LOGO EDL: coloque logo_EDL_oficial.png junto al script]")
        r.bold = True; r.font.color.rgb = BLANCO; r.font.size = Pt(11)

    doc.add_paragraph("")
    doc.add_paragraph("[TÍTULO]", style="Portada Título")
    doc.add_paragraph("[SUBTÍTULO]", style="Portada Subtítulo")
    doc.add_paragraph("")
    doc.add_paragraph("Autor: [AUTOR]", style="Portada Metadatos")
    doc.add_paragraph("Fecha: [FECHA]", style="Portada Metadatos")
    doc.add_paragraph("Versión: [VERSIÓN]", style="Portada Metadatos")

# ---------------------------------------------------------------- cuerpo
def _figura(doc, num):
    caja = doc.add_table(rows=1, cols=1)
    caja.style = doc.styles["Table Grid"]
    caja.alignment = 1
    c = caja.cell(0, 0)
    _anchos(caja, [ANCHO_UTIL_CM])
    fila = caja.rows[0]; fila.height = Cm(6); fila.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    pc = c.paragraphs[0]; pc.style = doc.styles["Tabla"]
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.add_run("[Espacio reservado para Figura %d: Insertar > Imagen. "
               "Sustituya este texto por la imagen real.]" % num)
    l1 = doc.add_paragraph(style="Leyenda")
    l1.add_run("Figura %d. [Título descriptivo de la figura]." % num)
    l2 = doc.add_paragraph(style="Leyenda")
    l2.add_run("Fuente: [FUENTE]  |  Periodo: [PERIODO]  |  Resultado: [RESULTADO]")

def _tabla_indicadores(doc):
    doc.add_paragraph("Tabla 1. [Descripción de los indicadores mostrados. "
                      "Complete con datos reales al editar.]", style="Leyenda")
    t = doc.add_table(rows=1, cols=4)
    t.style = doc.styles["Table Grid"]
    t.alignment = 1
    for i, nom in enumerate(("Indicador", "Valor", "Unidad", "Fuente")):
        c = t.rows[0].cells[i]
        _shd_celda(c, "F04A1D")
        pc = c.paragraphs[0]; pc.style = doc.styles["Tabla"]
        r = pc.add_run(nom); r.bold = True; r.font.color.rgb = BLANCO
    for k in range(1, 5):
        fila = t.add_row().cells
        vals = ["[Indicador %d]" % k, "[Valor]", "[Unidad]", "[Fuente]"]
        for i, v in enumerate(vals):
            pc = fila[i].paragraphs[0]; pc.style = doc.styles["Tabla"]
            pc.add_run(v)
    _anchos(t, [6.5, 3.0, 3.0, 3.5])
    return t

def cuerpo(doc):
    doc.add_paragraph("Resumen ejecutivo", style="Título 1")
    doc.add_paragraph("[Resumen ejecutivo: describa en 150-250 palabras el propósito, el "
                      "alcance y los hallazgos principales del análisis. Sustituya este campo "
                      "explicativo por contenido real; esta base no incluye estadísticas.]")

    doc.add_paragraph("Objetivos", style="Título 1")
    doc.add_paragraph("[Defina aquí el propósito general del informe.]")
    doc.add_paragraph("Objetivo general", style="Título 2")
    doc.add_paragraph("[Redacte el objetivo general del análisis.]")
    doc.add_paragraph("Objetivos específicos", style="Título 2")
    for i in (1, 2, 3):
        doc.add_paragraph("[Objetivo específico %d.]" % i, style="List Bullet")

    doc.add_paragraph("Datos y fuentes", style="Título 1")
    doc.add_paragraph("Conjuntos de datos", style="Título 2")
    doc.add_paragraph("[Describa cada conjunto de datos: nombre, cobertura, granularidad, "
                      "volumen aproximado y responsable.]")
    doc.add_paragraph("Fuentes y adquisición", style="Título 2")
    doc.add_paragraph("[Indique el origen de los datos (sistemas internos, APIs, encuestas, "
                      "fuentes públicas) y el procedimiento de extracción o recepción.]")

    doc.add_paragraph("Metodología", style="Título 1")
    doc.add_paragraph("[Explique el tratamiento aplicado: limpieza, validación, "
                      "transformaciones, técnicas de análisis y herramientas utilizadas.]")
    doc.add_paragraph("[Documente aquí los criterios de calidad de datos y las reglas de "
                      "negocio aplicadas a los indicadores.]")

    doc.add_paragraph("Resultados", style="Título 1")
    doc.add_paragraph("Indicadores principales", style="Título 2")
    _tabla_indicadores(doc)
    doc.add_paragraph("Evidencia gráfica", style="Título 2")
    _figura(doc, 1)
    _figura(doc, 2)

    doc.add_paragraph("Limitaciones", style="Título 1")
    doc.add_paragraph("[Describa las limitaciones conocidas: sesgos, vacíos de información, "
                      "supuestos metodológicos y alcance de las conclusiones.]")

    doc.add_paragraph("Conclusiones", style="Título 1")
    doc.add_paragraph("[Sintetice las conclusiones derivadas exclusivamente de los resultados "
                      "presentados, enlazadas con los objetivos.]")

    doc.add_paragraph("Referencias", style="Título 1")
    doc.add_paragraph("[Liste aquí las referencias reales consultadas, en formato APA u otro "
                      "estándar adoptado. Esta base no incluye bibliografía inventada.]")

# ---------------------------------------------------------------- encabezado / pie
def encabezado_y_pie(sec):
    h = sec.header; h.is_linked_to_previous = False
    ph = h.paragraphs[0]
    _borde_inferior(ph)                       # antes de alignment/spacing (orden OOXML)
    ph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rh = ph.add_run("Informe técnico de análisis de datos  ·  EDL Data Lab")
    rh.font.size = Pt(9); rh.font.color.rgb = TEXTO

    f = sec.footer; f.is_linked_to_previous = False
    pf = f.paragraphs[0]
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = pf.add_run("Página "); r0.font.size = Pt(9); r0.font.color.rgb = TEXTO
    _campo(pf, " PAGE ")
    r1 = pf.add_run(" de "); r1.font.size = Pt(9); r1.font.color.rgb = TEXTO
    _campo(pf, " NUMPAGES ")
    for r in pf.runs:
        r.font.size = Pt(9); r.font.color.rgb = TEXTO

# ---------------------------------------------------------------- verificación
def verificar(ruta, logo):
    d = Document(ruta)
    res = []
    def chk(desc, cond):
        res.append((desc, bool(cond)))
    nombres = {s.name for s in d.styles}
    for n in ("Normal", "Título 1", "Título 2", "Leyenda", "Tabla"):
        chk('Estilo real "%s" definido' % n, n in nombres)
    nm = d.styles["Normal"]
    chk("Normal: Arial 11 pt, #222222", nm.font.name == "Arial" and nm.font.size == Pt(11)
        and nm.font.color.rgb == TEXTO)
    chk("Normal: interlineado 1,15 y 6 pt después",
        abs(nm.paragraph_format.line_spacing - 1.15) < 1e-9
        and nm.paragraph_format.space_after == Pt(6))
    chk("Ocho secciones con Título 1",
        sum(p.style.name == "Título 1" for p in d.paragraphs) == 8)
    chk("Jerarquía Título 2 presente",
        sum(p.style.name == "Título 2" for p in d.paragraphs) >= 2)
    chk("Dos secciones (portada + cuerpo)", len(d.sections) == 2)
    s1, s2 = d.sections
    chk("Salto de sección a página nueva tras portada", s2.start_type == WD_SECTION.NEW_PAGE)
    chk("Portada sin encabezado propio (vacío)", not s1.header.paragraphs[0].text.strip())
    chk("Portada sin pie ni número (sin parte propia)",
        s1.footer.is_linked_to_previous and not s1.footer.paragraphs[0].text.strip())
    chk("Cuerpo: encabezado propio desvinculado",
        s2.header.is_linked_to_previous is False and s2.header.paragraphs[0].text.strip() != "")
    chk("Cuerpo: pie desvinculado con campo PAGE",
        s2.footer.is_linked_to_previous is False
        and any("PAGE" in p._p.xml for p in s2.footer.paragraphs))
    chk("Numeración continua (sin reinicio explícito)",
        "pgNumType" not in s2._sectPr.xml)
    tind = None
    for t in d.tables:
        if len(t.columns) == 4 and t.rows[0].cells[0].text.strip() == "Indicador":
            tind = t
    chk("Tabla editable 4 columnas (Indicador/Valor/Unidad/Fuente)",
        tind is not None and len(tind.rows) >= 2)
    chk("Dos espacios reservados para figuras",
        sum(1 for t in d.tables if len(t.columns) == 1
            and "Espacio reservado para Figura" in t.rows[0].cells[0].text) == 2)
    chk("Leyendas con campos [FUENTE] [PERIODO] [RESULTADO]",
        sum(1 for p in d.paragraphs if p.style.name == "Leyenda" and "[FUENTE]" in p.text
            and "[PERIODO]" in p.text and "[RESULTADO]" in p.text) == 2)
    chk("Banda negra #090909 en portada",
        any('w:fill="090909"' in t._tbl.xml for t in d.tables))
    chk("A4 y márgenes 2,5 cm en todas las secciones",
        all(s.page_width == Cm(21) and s.page_height == Cm(29.7)
            and s.left_margin == Cm(2.5) and s.top_margin == Cm(2.5) for s in d.sections))
    if logo and os.path.exists(logo):
        sh = d.inline_shapes
        chk("Logo insertado 1 vez, 12 cm de ancho (<= 16 cm útiles, sin recorte)",
            len(sh) == 1 and sh[0].width == Cm(12) and sh[0].height > 0)
    else:
        chk("Logo no encontrado: se insertó marcador de posición", True)
    print("\n--- Verificación automática de %s ---" % ruta)
    for desc, ok in res:
        print(" [%s] %s" % ("OK" if ok else "REVISAR", desc))
    print("Resultado global: %s" % ("TODO OK" if all(o for _, o in res) else "HAY PUNTOS A REVISAR"))
    return all(o for _, o in res)

# ---------------------------------------------------------------- main
def main():
    logo = sys.argv[1] if len(sys.argv) > 1 else "logo_EDL_oficial.png"
    doc = Document()
    configurar_estilos(doc)
    portada(doc, logo)
    doc.add_section(WD_SECTION.NEW_PAGE)      # salto real tras la portada
    cuerpo(doc)
    for s in doc.sections:
        configurar_pagina(s)
    encabezado_y_pie(doc.sections[1])         # solo cuerpo; portada queda sin número
    doc.save(SALIDA)
    print("Generado:", SALIDA)
    verificar(SALIDA, logo)

if __name__ == "__main__":
    main()